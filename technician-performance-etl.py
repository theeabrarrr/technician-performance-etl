from datetime import datetime
import os
from pathlib import Path
import shutil
import numpy as np
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
import pandas as pd

# ==========================================
# 1. DIRECTORY CONFIGURATION
# ==========================================
TEMP_DIR = Path(r"C:\temp")
BASE_DIR = Path(r"C:\Technician_Performance")
RAW_DIR = BASE_DIR / "raw_data"
REPORT_DIR = BASE_DIR / "reports"

RAW_DIR.mkdir(parents=True, exist_ok=True)
REPORT_DIR.mkdir(parents=True, exist_ok=True)


def get_latest_file(folder: Path, pattern: str) -> Path:
  """Finds the most recently modified file matching pattern."""
  files = list(folder.glob(pattern))
  return max(files, key=os.path.getmtime) if files else None


# ==========================================
# 2. AUTO-FETCH FILES FROM C:\temp
# ==========================================
latest_cancel = get_latest_file(TEMP_DIR, "CANCEL_NIL_TRANSFER*.xls*")
latest_feedback = get_latest_file(TEMP_DIR, "quality_feedback_report_*.csv")

if latest_cancel and latest_feedback:
  cancel_path = RAW_DIR / latest_cancel.name
  feedback_path = RAW_DIR / latest_feedback.name
  shutil.move(str(latest_cancel), str(cancel_path))
  shutil.move(str(latest_feedback), str(feedback_path))
  print(f"[INFO] Moved new ERP files from C:\\temp to {RAW_DIR}")
else:
  cancel_path = get_latest_file(RAW_DIR, "CANCEL_NIL_TRANSFER*.xls*")
  feedback_path = get_latest_file(RAW_DIR, "quality_feedback_report_*.csv")

if not cancel_path or not feedback_path:
  raise FileNotFoundError(
      "Could not find required files in C:\\temp or"
      " C:\\Technician_Performance\\raw_data."
  )

print(f"[PROCESSING] Cancel File:   {cancel_path.name}")
print(f"[PROCESSING] Feedback File: {feedback_path.name}")

# ==========================================
# 3. LOAD AND HARMONIZE DATA
# ==========================================
df_csv = pd.read_csv(feedback_path)
df_xls = pd.read_excel(cancel_path)

req_csv = [
    "BRANCH_NAME",
    "CUSTOMER_NAME",
    "PHONE_NO",
    "COMPLAINT_TYPE",
    "TECHNICIAN_NAME",
    "PURCHASE_DATE",
    "COMPLAINT_DATE",
    "COMPLAINT_NO",
    "MODEL_NAME",
    "SERIAL",
    "REMARKS",
    "COMPLETE_DATE",
    "COMPLETED_STATUS",
    "CLOSED_DATE",
    "DAYS",
]

req_xls = [
    "BRANCH_NAME",
    "CLOSED_DATE",
    "CUSTOMER_NAME",
    "COMPLAINT_DATE",
    "COMPLAINT_NO",
    "MODEL_NAME",
    "TECHNICIAN_NAME",
    "SERIAL",
    "STATUS",
    "REMARKS",
    "COMPLAINT_TYPE",
    "BRANCH_REMARKS",
]

df_csv_sub = df_csv[[c for c in req_csv if c in df_csv.columns]].copy()
df_csv_sub.rename(columns={"COMPLETED_STATUS": "STATUS"}, inplace=True)
df_xls_sub = df_xls[[c for c in req_xls if c in df_xls.columns]].copy()

for col in df_xls_sub.columns:
  if col not in df_csv_sub.columns:
    df_csv_sub[col] = np.nan

for col in df_csv_sub.columns:
  if col not in df_xls_sub.columns:
    df_xls_sub[col] = np.nan

csv_date_clean = (
    df_csv_sub["CLOSED_DATE"].astype(str).str.replace("Sept", "Sep", regex=False)
)
df_csv_sub["_SORT_DATE"] = pd.to_datetime(
    csv_date_clean, format="%d-%b-%y", errors="coerce"
)
df_xls_sub["_SORT_DATE"] = pd.to_datetime(
    df_xls_sub["CLOSED_DATE"], errors="coerce"
)

# Priority: COMPLETED takes precedence over NIL when dates tie
df_csv_sub["_PRIORITY"] = 2
df_xls_sub["_PRIORITY"] = 1

combined = pd.concat([df_csv_sub, df_xls_sub], ignore_index=True)
combined.sort_values(
    by=["COMPLAINT_NO", "_SORT_DATE", "_PRIORITY"],
    ascending=[True, True, True],
    inplace=True,
)

master_df = combined.drop_duplicates(
    subset=["COMPLAINT_NO"], keep="last"
).copy()
master_df.drop(columns=["_SORT_DATE", "_PRIORITY"], inplace=True)

column_order = [
    "BRANCH_NAME",
    "TECHNICIAN_NAME",
    "COMPLAINT_NO",
    "STATUS",
    "COMPLAINT_TYPE",
    "CUSTOMER_NAME",
    "PHONE_NO",
    "MODEL_NAME",
    "SERIAL",
    "COMPLAINT_DATE",
    "COMPLETE_DATE",
    "CLOSED_DATE",
    "DAYS",
    "REMARKS",
    "BRANCH_REMARKS",
    "PURCHASE_DATE",
]
master_df = master_df.reindex(
    columns=[c for c in column_order if c in master_df.columns]
)

for d_col in ["COMPLAINT_DATE", "COMPLETE_DATE", "CLOSED_DATE", "PURCHASE_DATE"]:
  if d_col in master_df.columns:
    clean_s = (
        master_df[d_col].astype(str).str.replace("Sept", "Sep", regex=False)
    )
    dt_parsed = pd.to_datetime(clean_s, format="mixed", errors="coerce")
    master_df[d_col] = dt_parsed.dt.strftime("%Y-%m-%d").fillna("")

# Object casting to prevent LossySetitemError in modern Pandas
master_df = master_df.astype(object).where(pd.notnull(master_df), None)

# ==========================================
# 4. FILTER OUT TRANSFERED FOR ACCURATE KPI
# ==========================================
# Transferred jobs belong to other zones/centers and are excluded from technician evaluation
perf_df = master_df[master_df["STATUS"] != "TRANSFERED"].copy()

pvt = pd.pivot_table(
    perf_df,
    index="TECHNICIAN_NAME",
    columns="STATUS",
    values="COMPLAINT_NO",
    aggfunc="count",
    fill_value=0,
)

expected_statuses = ["COMPLETED", "CANCELED", "REJECTED", "NIL"]
for st in expected_statuses:
  if st not in pvt.columns:
    pvt[st] = 0

status_cols = [c for c in expected_statuses if c in pvt.columns]
pvt = pvt[status_cols]
pvt["Total Assigned"] = pvt.sum(axis=1)
pvt["Completion Rate (%)"] = (
    (pvt["COMPLETED"] / pvt["Total Assigned"]) * 100
).round(2)
pvt.sort_values(by="Total Assigned", ascending=False, inplace=True)
pvt.reset_index(inplace=True)

# ==========================================
# 5. OPENPYXL STYLING & EMBEDDED VISUALS
# ==========================================
timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
output_file = REPORT_DIR / f"Technician_Performance_Report_{timestamp_str}.xlsx"

wb = openpyxl.Workbook()
ws_perf = wb.active
ws_perf.title = "Technician_Performance"
ws_data = wb.create_sheet(title="Consolidated_Data")

NAVY_HEADER_FILL = PatternFill(
    start_color="1F4E78", end_color="1F4E78", fill_type="solid"
)
KPI_BG_FILL = PatternFill(
    start_color="F2F4F7", end_color="F2F4F7", fill_type="solid"
)
ZEBRA_FILL = PatternFill(
    start_color="F9FAFB", end_color="F9FAFB", fill_type="solid"
)
TOTAL_ROW_FILL = PatternFill(
    start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
)

FONT_TITLE = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
FONT_SUBTITLE = Font(name="Calibri", size=10, italic=True, color="D9D9D9")
FONT_TH = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_BODY = Font(name="Calibri", size=10, color="000000")
FONT_TOTAL = Font(name="Calibri", size=11, bold=True, color="000000")
FONT_KPI_VAL = Font(name="Calibri", size=16, bold=True, color="1F4E78")
FONT_KPI_LBL = Font(name="Calibri", size=9, bold=True, color="595959")

BORDER_THIN = Side(border_style="thin", color="D9D9D9")
BORDER_DOUBLE = Side(border_style="double", color="1F4E78")
CELL_BORDER = Border(
    left=BORDER_THIN, right=BORDER_THIN, top=BORDER_THIN, bottom=BORDER_THIN
)
TOTAL_BORDER = Border(
    top=BORDER_THIN, bottom=BORDER_DOUBLE, left=BORDER_THIN, right=BORDER_THIN
)

# Header Banner
ws_perf.merge_cells("A1:F1")
ws_perf["A1"] = "SERVICE CENTER - TECHNICIAN PERFORMANCE DASHBOARD"
ws_perf["A1"].font = FONT_TITLE
ws_perf["A1"].fill = NAVY_HEADER_FILL
ws_perf["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws_perf.row_dimensions[1].height = 30

ws_perf.merge_cells("A2:F2")
ws_perf["A2"] = (
    f"Report Generated on: {datetime.now().strftime('%d-%b-%Y %I:%M %p')} |"
    " Note: Transferred Out-of-Zone Calls Excluded"
)
ws_perf["A2"].font = FONT_SUBTITLE
ws_perf["A2"].fill = NAVY_HEADER_FILL
ws_perf["A2"].alignment = Alignment(horizontal="center", vertical="center")
ws_perf.row_dimensions[2].height = 18

# Top KPI Summary Cards (Calculated purely on active zone work)
tot_actionable = len(perf_df)
tot_completed = int((perf_df["STATUS"] == "COMPLETED").sum())
eff_rate = (tot_completed / max(tot_actionable, 1)) * 100

kpis = [
    ("ASSIGNED IN-ZONE", tot_actionable, "A", "B"),
    ("TOTAL COMPLETED", tot_completed, "C", "D"),
    ("ZONE EFFICIENCY", f"{eff_rate:.1f}%", "E", "F"),
]

for title, val, start_c, end_c in kpis:
  ws_perf.merge_cells(f"{start_c}4:{end_c}4")
  ws_perf.merge_cells(f"{start_c}5:{end_c}5")
  lbl_cell = ws_perf[f"{start_c}4"]
  val_cell = ws_perf[f"{start_c}5"]
  lbl_cell.value = title
  lbl_cell.font = FONT_KPI_LBL
  lbl_cell.fill = KPI_BG_FILL
  lbl_cell.alignment = Alignment(horizontal="center", vertical="center")
  val_cell.value = val
  val_cell.font = FONT_KPI_VAL
  val_cell.fill = KPI_BG_FILL
  val_cell.alignment = Alignment(horizontal="center", vertical="center")
  for col_idx in [start_c, end_c]:
    ws_perf[f"{col_idx}4"].border = CELL_BORDER
    ws_perf[f"{col_idx}5"].border = CELL_BORDER

ws_perf.row_dimensions[4].height = 18
ws_perf.row_dimensions[5].height = 25

# Performance Table Headers
table_start_row = 7
headers = list(pvt.columns)
for col_idx, h_text in enumerate(headers, 1):
  c = ws_perf.cell(row=table_start_row, column=col_idx, value=h_text)
  c.font = FONT_TH
  c.fill = NAVY_HEADER_FILL
  c.alignment = Alignment(
      horizontal="center" if col_idx > 1 else "left", vertical="center"
  )
  c.border = CELL_BORDER
ws_perf.row_dimensions[table_start_row].height = 24

# Data Rows
current_row = table_start_row + 1
for _, row_data in pvt.iterrows():
  for col_idx, val in enumerate(row_data, 1):
    c = ws_perf.cell(row=current_row, column=col_idx, value=val)
    c.font = FONT_BODY
    c.border = CELL_BORDER
    if col_idx == 1:
      c.alignment = Alignment(horizontal="left", vertical="center")
    elif col_idx == len(row_data):
      c.alignment = Alignment(horizontal="right", vertical="center")
      c.number_format = '0.0"%"'
    else:
      c.alignment = Alignment(horizontal="right", vertical="center")
      c.number_format = "#,##0"
    if (current_row % 2) == 0:
      c.fill = ZEBRA_FILL
  ws_perf.row_dimensions[current_row].height = 19
  current_row += 1

last_data_row = current_row - 1

# Dynamic Grand Total Row
ws_perf.cell(row=current_row, column=1, value="Grand Total").font = FONT_TOTAL
ws_perf.cell(row=current_row, column=1).fill = TOTAL_ROW_FILL
ws_perf.cell(row=current_row, column=1).border = TOTAL_BORDER
ws_perf.cell(row=current_row, column=1).alignment = Alignment(
    horizontal="left", vertical="center"
)

for col_idx in range(2, len(headers)):
  col_let = get_column_letter(col_idx)
  c = ws_perf.cell(
      row=current_row,
      column=col_idx,
      value=f"=SUM({col_let}8:{col_let}{last_data_row})",
  )
  c.font = FONT_TOTAL
  c.fill = TOTAL_ROW_FILL
  c.border = TOTAL_BORDER
  c.alignment = Alignment(horizontal="right", vertical="center")
  c.number_format = "#,##0"

comp_col = get_column_letter(headers.index("COMPLETED") + 1)
tot_col = get_column_letter(headers.index("Total Assigned") + 1)
rate_col_idx = len(headers)
rate_cell = ws_perf.cell(
    row=current_row,
    column=rate_col_idx,
    value=f"=({comp_col}{current_row}/{tot_col}{current_row})*100",
)
rate_cell.font = FONT_TOTAL
rate_cell.fill = TOTAL_ROW_FILL
rate_cell.border = TOTAL_BORDER
rate_cell.alignment = Alignment(horizontal="right", vertical="center")
rate_cell.number_format = '0.0"%"'
ws_perf.row_dimensions[current_row].height = 22

for col in ws_perf.columns:
  max_len = max(len(str(cell.value or "")) for cell in col)
  col_letter = get_column_letter(col[0].column)
  ws_perf.column_dimensions[col_letter].width = max(max_len + 4, 13)
ws_perf.column_dimensions["A"].width = 25

# ----------------------------------------------------
# EMBEDDED HORIZONTAL BAR CHART (NO TRANSFERED)
# ----------------------------------------------------
chart = BarChart()
chart.type = "bar"
chart.grouping = "stacked"
chart.overlap = 100
chart.title = "Technician Job Volume & Status Breakdown (Active Work)"
chart.style = 10
chart.height = 25
chart.width = 18
chart.y_axis.title = "Technician"
chart.x_axis.title = "Number of Complaints"

# Map to COMPLETED, CANCELED, REJECTED, NIL (Columns 2 to 5)
data = Reference(
    ws_perf, min_col=2, min_row=7, max_col=5, max_row=last_data_row
)
cats = Reference(ws_perf, min_col=1, min_row=8, max_row=last_data_row)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)

# Position chart cleanly next to the table
ws_perf.add_chart(chart, "H4")

# ----------------------------------------------------
# TAB 2: MASTER CONSOLIDATED DATA (PRESERVES ALL)
# ----------------------------------------------------
data_headers = list(master_df.columns)
ws_data.append(data_headers)

for row in master_df.itertuples(index=False):
  ws_data.append(list(row))

table_ref = f"A1:{get_column_letter(len(data_headers))}{len(master_df)+1}"
tab_style = Table(displayName="MasterComplaintsTable", ref=table_ref)
tab_style.tableStyleInfo = TableStyleInfo(
    name="TableStyleMedium9",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)
ws_data.add_table(tab_style)

for col in ws_data.columns:
  max_len = max(len(str(cell.value or "")) for cell in col)
  col_letter = get_column_letter(col[0].column)
  ws_data.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

wb.save(output_file)
print(
    f"\n[SUCCESS] Report exported cleanly without Transferred jobs to:\n{output_file}"
)